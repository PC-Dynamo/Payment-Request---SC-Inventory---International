#!/usr/bin/env python3
"""
seed_from_reference.py - one-off import of the team's payment workbook.

Reads the "updated" sheet of
  Reference/Payment Request for Four-week Rolling Forecast (*.xlsx)
and records, for the POs this app lists:

  * which deposits / balances are already PAID or REQUESTED (with the paid
    date, invoice number, amount and comment the team entered), and
  * the recent manual charge lines - freight (USD) and the AUD document /
    handling charges (customs brokerage, GST, duty, cartage) - as manual lines.

Idempotent and non-destructive: a line anyone has already touched in the app
is never overwritten, and a charge row is imported once (keyed by its content).

  python seed_from_reference.py [--dry-run] [--since 2026-08-01]
"""
from __future__ import annotations

import argparse
import datetime as dt
import hashlib
import json
import re
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import ledger as L  # noqa: E402
import payments as P  # noqa: E402

REF_DIR = HERE.parent / "Reference"
PAYLOAD = HERE / "report_payload.json"

CHARGE_RULES = [
    (r"customs\s*\+\s*gst|customs brokerage|brokerage", "Customs Brokerage"),
    (r"freight", "Freight Payment"),
    (r"\bgst\b", "GST Payment"),
    (r"duty", "Government Duty"),
    (r"local cartage", "Local Cartage Delivery"),
    (r"cartage", "Cartage"),
    (r"detention", "Detention Charges"),
    (r"inspection|qima", "Inspection"),
    (r"fob fee|exw|document|handling|storage", "Other"),
]


def _date(v) -> str | None:
    if v is None or v == "":
        return None
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    m = re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{4})$", s)
    if m:
        try:
            return dt.date(int(m[3]), int(m[2]), int(m[1])).isoformat()
        except ValueError:
            return None
    try:
        return dt.date.fromisoformat(s[:10]).isoformat()
    except ValueError:
        return None


def _num(v) -> float | None:
    if v in (None, ""):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"-?[\d,]+(?:\.\d+)?", str(v))
    try:
        return float(m.group(0).replace(",", "")) if m else None
    except ValueError:
        return None


def _txt(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").replace("\xa0", " ")).strip()


def _charge_type(status: str) -> str | None:
    s = status.lower()
    for rx, label in CHARGE_RULES:
        if re.search(rx, s):
            return label
    return None


# The team's "not finalised yet" colour: Excel theme Accent 2, lighter 80% (peach).
PEACH = {("theme", 5)}


def peach_rows(path: Path) -> set[int]:
    """Sheet row numbers the team shaded peach = amount / PI not finalised."""
    import openpyxl
    wb = openpyxl.load_workbook(path)        # styles need a full (not read-only) load
    ws = wb["updated"]
    out = set()
    for r in range(2, ws.max_row + 1):
        fl = ws.cell(row=r, column=4).fill
        if fl is None or fl.fill_type != "solid":
            continue
        c = fl.fgColor
        if (c.type == "theme" and (c.type, c.theme) in PEACH and (c.tint or 0) > 0.5) or \
                (c.type == "rgb" and str(c.rgb).upper() in ("FFFCE4D6", "FFF8CBAD", "FFFBE2D5")):
            out.add(r)
    wb.close()
    return out


def load_reference(path: Path) -> list[dict]:
    import openpyxl
    peach = peach_rows(path)
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb["updated"]
    it = ws.iter_rows(values_only=True)
    header = [_txt(h).lower() for h in next(it)]

    def col(name):
        for i, h in enumerate(header):
            if h.startswith(name):
                return i
        raise KeyError(name)

    ix = {k: col(n) for k, n in {
        "entry": "entry date", "vendor": "vendor entity", "supplier": "supplier name",
        "po": "po number", "doc": "document", "branch": "branch", "stage": "stage",
        "status": "payment status", "req": "payment request date", "due": "due date",
        "usd": "value (usd)", "aud": "value (aud)", "gbp": "value (gbp)", "eur": "value (eur)",
        "paid": "paid date", "paid_amt": "paid amount", "comment": "comment",
        "acc": "accounts scheduled"}.items()}
    rows = []
    for n, r in enumerate(it, start=2):
        if not r or not any(r):
            continue
        g = lambda k: r[ix[k]] if ix[k] < len(r) else None  # noqa: E731
        rows.append({
            "row": n, "entry": _date(g("entry")), "vendor": _txt(g("vendor")),
            "supplier": _txt(g("supplier")), "po_text": _txt(g("po")),
            "pos": P.po_list(str(g("po") or "")), "doc": _txt(g("doc")),
            "branch": _txt(g("branch")), "status": str(g("status") or "").strip(),
            "req": _date(g("req")), "due": _date(g("due")),
            "usd": _num(g("usd")), "aud": _num(g("aud")), "gbp": _num(g("gbp")),
            "eur": _num(g("eur")), "paid": _date(g("paid")),
            "paid_amt": _num(g("paid_amt")), "comment": _txt(g("comment")),
            "acc": _date(g("acc")),
            "peach": n in peach,
        })
    wb.close()
    return rows


def _money(row: dict) -> tuple[float | None, str | None]:
    for cur in ("usd", "aud", "gbp", "eur"):
        if row[cur]:
            return row[cur], cur.upper()
    return None, None


def plan(rows: list[dict], auto: list[dict], since: str, today: dt.date) -> tuple[dict, list]:
    kinds: dict[str, set] = {}
    for ln in auto:
        kinds.setdefault(ln["po"], set()).add(ln["kind"])
    states: dict[str, dict] = {}          # line_id -> values (later rows win)
    charges: list[dict] = []

    for row in sorted(rows, key=lambda x: (x["req"] or x["entry"] or "", x["row"])):
        status = row["status"]
        if not status:
            continue
        ctype = _charge_type(status)
        if ctype:
            if (row["entry"] or "") < since and (row["req"] or "") < since:
                continue
            if not row["pos"] or not any(p in kinds for p in row["pos"]):
                if (row["req"] or "") < today.isoformat() and row["paid"]:
                    continue       # paid charge for POs no longer open: history only
            amt, cur = _money(row)
            if amt is None:
                continue
            st = "Paid" if row["paid"] else ("Requested" if row["req"] and
                                             row["req"] <= today.isoformat() else None)
            charges.append({
                "seed_key": hashlib.sha1(json.dumps(
                    [row["po_text"], row["doc"], status, amt, row["req"]]).encode()).hexdigest(),
                "entry_date": row["entry"], "vendor": row["vendor"] or "",
                "pay_type": ctype, "po_refs": ", ".join(row["pos"]) or row["po_text"][:200],
                "branch": row["branch"], "doc_no": row["doc"], "currency": cur,
                "amount": amt, "due_date": row["due"], "request_date": row["req"],
                "status": st, "paid_date": row["paid"],
                "comment": (row["comment"] or f"{status} (imported)")[:600],
                "acc_scheduled_date": row["acc"], "paid_amount": row["paid_amt"],
                "finalised": None if row["paid"] else (0 if row["peach"] else 1),
            })
            continue

        # ---- supplier payment row: only for a single PO this app lists
        if len(row["pos"]) != 1 or row["pos"][0] not in kinds:
            continue
        po = row["pos"][0]
        k = kinds[po]
        final = "balance" if "balance" in k else "full"
        amt, cur = _money(row)
        if amt is None and row["paid_amt"]:
            amt = row["paid_amt"]
            cur = "GBP" if "gbp" in str(row.get("comment", "")).lower() else cur
        for piece in [p.strip().lower() for p in status.splitlines() if p.strip()]:
            is_for = bool(re.search(r"\bfor\b", piece))
            paid = ("paid" in piece and not is_for) or bool(row["paid"] and len(
                [p for p in status.splitlines() if p.strip()]) == 1)
            if piece.startswith("allocated from"):
                targets, val = [final], {"status": "Not Payable",
                                         "comment": f"{status.strip()} (per payment workbook)"}
            elif "deposit" in piece or ("advance" in piece and "100%" not in piece) \
                    or "tooling" in piece:
                targets = ["deposit"] if "deposit" in k else [final]
                val = {"status": "Paid" if paid else "Requested"}
            elif "balance" in piece:
                targets, val = [final], {"status": "Paid" if paid else "Requested"}
            elif paid:
                # "Paid" / "100% paid" / "100% T/T paid": the whole PO is settled.
                targets, val = sorted(k), {"status": "Paid"}
            elif is_for or "payment" in piece or "advance" in piece or "t/t" in piece:
                targets, val = [final], {"status": "Requested"}
            else:
                continue
            for kind in targets:
                lid = f"{po}|{kind}"
                v = dict(val)
                if v["status"] == "Requested" and not (row["req"] and row["req"] <= today.isoformat()):
                    # Queued in the workbook but not yet sent: leave the status to
                    # the app, keep the planned request date, amount and invoice.
                    v["status"] = None
                if v["status"] in ("Paid", "Requested", None):
                    v["request_date"] = row["req"]
                    if v["status"] == "Paid":
                        v["paid_date"] = row["paid"] or row["req"]
                    if amt and len(targets) == 1:
                        v["amount"], v["currency"] = amt, cur
                v["invoice_no"] = row["doc"]
                v["vendor_entity"] = row["vendor"] or None
                if v.get("status") != "Paid":
                    v["finalised"] = 0 if row["peach"] else 1
                v["acc_scheduled_date"] = row["acc"]
                if v.get("status") == "Paid" and row["paid_amt"]:
                    v["paid_amount"] = row["paid_amt"]
                if row["comment"] and "comment" not in v:
                    v["comment"] = row["comment"]
                v["_row"] = row["row"]
                states[lid] = v
    return states, charges


NEW_LINE_FIELDS = ("vendor_entity", "acc_scheduled_date", "paid_amount")


def backfill(states: dict, charges: list) -> int:
    n = 0
    with L.connect() as con:
        for lid, v in states.items():
            cur = con.execute("SELECT * FROM line_state WHERE line_id = ?", (lid,)).fetchone()
            if not cur or not str(cur["updated_by"] or "").startswith("seed"):
                continue
            for f in NEW_LINE_FIELDS:
                val = L.clean(f, v.get(f))
                if val not in (None, "") and cur[f] in (None, ""):
                    con.execute(f"UPDATE line_state SET {f} = ? WHERE line_id = ?", (val, lid))
                    n += 1
        for c in charges:
            cur = con.execute("SELECT * FROM manual_line WHERE seed_key = ?", (c["seed_key"],)).fetchone()
            if not cur or cur["updated_by"]:        # edited in the app since: leave it
                continue
            for f in ("acc_scheduled_date", "paid_amount"):
                val = L.clean(f, c.get(f))
                if val not in (None, "") and cur[f] in (None, ""):
                    con.execute(f"UPDATE manual_line SET {f} = ? WHERE id = ?", (val, cur["id"]))
                    n += 1
    return n


SYNC_LINE_FIELDS = ("status", "request_date", "paid_date", "amount", "currency", "invoice_no",
                    "comment", "vendor_entity", "acc_scheduled_date", "paid_amount", "finalised")
SYNC_MANUAL_FIELDS = ("status", "request_date", "paid_date", "due_date", "amount", "currency",
                      "doc_no", "comment", "acc_scheduled_date", "paid_amount", "vendor", "finalised")
SEED_WHO = "seed: payment request workbook"


def _is_seed(who) -> bool:
    return not who or str(who).startswith("seed")


def sync(states: dict, charges: list, label: str) -> dict:
    """Bring the app up to a newer copy of the team's workbook.

    Lines still holding only workbook data follow the workbook; new rows are
    added; anything a person changed in the app is kept and reported."""
    out = {"line_added": 0, "line_updated": 0, "manual_added": 0, "manual_updated": 0,
           "conflicts": []}
    now = L.now_awst()
    with L.connect() as con:
        for lid, v in states.items():
            vals = {f: L.clean(f, v.get(f)) for f in SYNC_LINE_FIELDS if f in v}
            cur = con.execute("SELECT * FROM line_state WHERE line_id = ?", (lid,)).fetchone()
            if not cur:
                cols = ["line_id", *vals, "seed_key", "updated_by", "updated_at"]
                con.execute(f"INSERT INTO line_state({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
                            [lid, *vals.values(), f"reference row {v.get('_row')}", SEED_WHO, now])
                out["line_added"] += 1
                continue
            diff = {f: x for f, x in vals.items() if x is not None and x != cur[f]}
            if v.get("status", "-") is None and cur["status"] and _is_seed(cur["updated_by"]):
                diff["status"] = None        # back to automatic: queued, not yet requested
            if not diff:
                continue
            if not _is_seed(cur["updated_by"]):
                out["conflicts"].append(f"{lid}: kept the app's {', '.join(sorted(diff))} "
                                        f"(edited by {cur['updated_by']})")
                continue
            sets = ", ".join(f"{f} = ?" for f in diff)
            con.execute(f"UPDATE line_state SET {sets}, updated_by = ?, updated_at = ? WHERE line_id = ?",
                        [*diff.values(), SEED_WHO, now, lid])
            con.execute("INSERT INTO audit(at, who, action, line_id, detail) VALUES (?,?,?,?,?)",
                        (now, SEED_WHO, f"sync from {label}", lid, json.dumps(diff, default=str)))
            out["line_updated"] += 1

        for c in charges:
            key = c["seed_key"]
            cur = con.execute("SELECT * FROM manual_line WHERE seed_key = ? AND deleted = 0", (key,)).fetchone()
            if not cur and c.get("doc_no"):
                # Same invoice, re-dated or re-valued in the workbook: update, do not duplicate.
                cur = con.execute(
                    "SELECT * FROM manual_line WHERE doc_no = ? AND pay_type = ? AND deleted = 0 "
                    "AND created_by LIKE 'seed%'", (c["doc_no"], c["pay_type"])).fetchone()
            if not cur:
                data = {k: x for k, x in c.items() if k != "seed_key"}
                vals = {f: L.clean(f, data.get(f)) for f in L.MANUAL_FIELDS}
                vals["pay_type"] = vals["pay_type"] or "Other"
                if vals["amount"] is None:
                    continue
                vals["currency"] = vals["currency"] or L.CHARGE_TYPES.get(vals["pay_type"], "AUD")
                cols = ["seed_key", *L.MANUAL_FIELDS, "created_by", "created_at"]
                con.execute(f"INSERT INTO manual_line({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
                            [key, *[vals[f] for f in L.MANUAL_FIELDS], SEED_WHO, now])
                out["manual_added"] += 1
                continue
            vals = {f: L.clean(f, c.get(f)) for f in SYNC_MANUAL_FIELDS if f in c}
            diff = {f: x for f, x in vals.items() if x not in (None, "") and x != cur[f]}
            if not diff:
                continue
            if not _is_seed(cur["updated_by"]):
                out["conflicts"].append(f"M{cur['id']} {cur['pay_type']} {cur['doc_no'] or ''}: kept the "
                                        f"app's {', '.join(sorted(diff))} (edited by {cur['updated_by']})")
                continue
            sets = ", ".join(f"{f} = ?" for f in diff)
            con.execute(f"UPDATE manual_line SET {sets}, updated_at = ? WHERE id = ?",
                        [*diff.values(), now, cur["id"]])
            con.execute("INSERT INTO audit(at, who, action, line_id, detail) VALUES (?,?,?,?,?)",
                        (now, SEED_WHO, f"sync from {label}", f"M{cur['id']}", json.dumps(diff, default=str)))
            out["manual_updated"] += 1
    return out


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--since", default="2026-08-01",
                    help="import charge lines entered or requested on/after this date")
    ap.add_argument("--file", help="reference workbook (default: newest in Reference/)")
    ap.add_argument("--sync", action="store_true",
                    help="bring the app up to a newer copy of the workbook (keeps people's edits)")
    ap.add_argument("--backfill", action="store_true",
                    help="only fill the register columns added later (vendor entity, accounts "
                         "scheduled date, paid amount) on lines the seed created and nobody edited")
    a = ap.parse_args()

    ref = Path(a.file) if a.file else max(REF_DIR.glob("Payment Request*.xlsx"),
                                          key=lambda p: p.stat().st_mtime)
    rows = load_reference(ref)
    auto = json.loads(PAYLOAD.read_text(encoding="utf-8"))["auto_lines"]
    today = dt.date.today()
    states, charges = plan(rows, auto, a.since, today)
    print(f"reference: {ref.name} - {len(rows)} rows")
    print(f"plan: {len(states)} auto-line states, {len(charges)} manual charge lines")
    if a.sync:
        res = sync(states, charges, ref.name)
        print(f"sync: {res['line_added']} lines added, {res['line_updated']} updated; "
              f"{res['manual_added']} manual lines added, {res['manual_updated']} updated; "
              f"{len(res['conflicts'])} kept as edited in the app")
        for x in res["conflicts"][:40]:
            print("  kept:", x)
        return 0
    if a.backfill:
        n = backfill(states, charges)
        print(f"backfilled {n} empty register fields")
        return 0
    if a.dry_run:
        for lid, v in list(states.items())[:40]:
            print("  ", lid, {k: v[k] for k in v if k != "_row"})
        for c in charges[:40]:
            print("  +", c["pay_type"], c["po_refs"][:40], c["currency"], c["amount"],
                  c["status"], c["request_date"])
        return 0
    n_state = sum(L.seed_line_state(lid, {k: v for k, v in val.items() if k != "_row"},
                                    f"reference row {val['_row']}")
                  for lid, val in states.items())
    n_charge = 0
    for c in charges:
        key = c.pop("seed_key")
        if L.add_manual(c, "seed: payment request workbook", seed_key=key):
            n_charge += 1
    print(f"seeded {n_state} line states (others already recorded), "
          f"{n_charge} new manual lines")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
