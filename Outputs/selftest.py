#!/usr/bin/env python3
"""
selftest.py - checks the rules and the built data. Exit 1 on any failure.

  python selftest.py          # rules + checks on the current cache/payload
"""
from __future__ import annotations

import datetime as dt
import json
import sys
import tempfile
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import ledger as L  # noqa: E402
import payments as P  # noqa: E402
import terms as T  # noqa: E402

FAILS: list[str] = []


def check(cond, msg):
    if not cond:
        FAILS.append(msg)
        print("  FAIL", msg)


def D(s):
    return dt.date.fromisoformat(s)


def test_request_days():
    # USD runs Tuesday, AUD runs Thursday. 2026-09-25 Fri, 09-29 Tue, 10-01 Thu.
    tue, thu = (1,), (3,)
    check(P.next_request_day(D("2026-09-29"), tue) == D("2026-09-29"), "Tuesday is a USD pay run")
    check(P.next_request_day(D("2026-09-25"), tue) == D("2026-09-29"), "Friday -> next Tuesday")
    check(P.next_request_day(D("2026-09-29"), thu) == D("2026-10-01"), "AUD charges -> Thursday")
    check(P.request_day_for(D("2026-10-08"), D("2026-09-25"), tue) == D("2026-10-06"),
          "due Thu 8 Oct -> Tuesday 6 Oct run")
    check(P.request_day_for(D("2026-10-13"), D("2026-09-25"), tue) == D("2026-10-13"),
          "due on a Tuesday -> that Tuesday")
    check(P.request_day_for(D("2026-01-01"), D("2026-09-25"), tue) == D("2026-09-29"),
          "overdue -> next pay run, never dropped")
    check(P.weekdays_for("AUD") == (3,) and P.weekdays_for("USD") == (1,) and P.weekdays_for("GBP") == (1,),
          "pay-run day follows the currency")


def test_po_list():
    check(P.po_list("PO 128830 - Impulse \nPO-127734 / PO128830") == ["PO-128830", "PO-127734"],
          "po_list parses and de-duplicates mixed PO spellings")
    check(P.po_list("PI-2026-1") == [], "po_list ignores short numbers")


def test_terms():
    cases = [
        ("Full payment - should be made at least one week before the shipment arrives", None, 0, "eta", -7),
        ("20% Deposit, 80% against copy of shipping docs", "20", 20, "etd", 7),
        ("45 days from the sailing day", None, 0, "etd", 45),
        ("100% payment within 14 days of container receipt", "100", 0, "eta", 14),
        ("30% deposit,  Balance payment when products arrieved Lisbon, Portugal port.", "30", 30, "eta", 0),
        ("100% advance payment", "100", 0, "order", 0),
        ("30% T/T before mass production,70% at sight at B/L", None, 30, "etd", 7),
        ("40% Deposit, 60% to be paid before booking the shipping space", "40", 40, "etd", -7),
        ("30% Deposit, 70% final payment - Paid 7 days after shipment", "30", 30, "etd", 7),
    ]
    for text, dep, want_dep, anchor, off in cases:
        r = T.parse_terms({"terms": text, "deposit_pct": dep,
                           "due_days": 45 if "sailing" in text else (14 if "14 days" in text else
                                                                     (7 if "7 days after" in text else None)),
                           "basis": ""})
        check((r["deposit_pct"], r["balance_anchor"], r["balance_offset"]) == (want_dep, anchor, off),
              f"terms {text[:40]!r} -> {r['deposit_pct'], r['balance_anchor'], r['balance_offset']}")
    check(T.parse_terms(None)["source"] == "default", "missing terms fall back to a default")

    book = T.TermsBook([{"supplier_name": "Nantong Kylin"}, {"supplier_name": "Qingdao CSP Industry and Trade Co.,ltd"},
                        {"supplier_name": "Shandong VOG Sports Products Co Ltd"}, {"supplier_name": "FDF"}])
    check(book.match("NANTONG RUIMING SUPPLY CHAIN CO.,LTD.") is None, "a shared city name is not a match")
    check(book.match("SHANDONG CROSSMAX FITNESS CO.,LTD") is None, "Crossmax is not VOG")
    check((book.match("FDF Limited") or {}).get("supplier_name") == "FDF", "FDF Limited -> FDF")
    check((book.match("Qingdao CSP Industry and Trade Co.,ltd") or {}).get("supplier_name")
          == "Qingdao CSP Industry and Trade Co.,ltd", "exact name matches")


def test_merge():
    base = {"source": "auto", "po": "PO-10001", "po_list": ["PO-10001"], "vendor": "X", "state": "WA",
            "segment": "Retail", "currency": "USD", "flags": [], "assumed_paid": None}
    auto = [
        {**base, "line_id": "PO-10001|deposit", "kind": "deposit", "amount_auto": 30.0, "due_auto": "2026-09-01"},
        {**base, "line_id": "PO-10001|balance", "kind": "balance", "amount_auto": 70.0, "due_auto": "2026-11-20"},
        {**base, "line_id": "PO-2|full", "po": "PO-2", "kind": "full", "amount_auto": 5.0, "due_auto": None},
        {**base, "line_id": "PO-3|deposit", "po": "PO-3", "kind": "deposit", "amount_auto": 9.0,
         "due_auto": "2026-01-01", "assumed_paid": "deposit assumed paid"},
    ]
    states = {"PO-10001|balance": {"amount": 71.5, "status": None}}
    manual = [{"id": 1, "po_refs": "PO 10001", "pay_type": "GST Payment", "currency": "AUD", "amount": 10.0,
               "due_date": None, "status": None}]
    rows, when = P.merge(auto, states, manual, D("2026-09-25"), 4)
    by = {r["line_id"]: r for r in rows}
    check(by["PO-10001|deposit"]["status"] == "For Payment" and by["PO-10001|deposit"]["overdue"],
          "overdue deposit is For Payment and flagged")
    check(by["PO-10001|deposit"]["request_date"] == "2026-09-29", "overdue line joins the next Tuesday run")
    check(by["PO-10001|balance"]["status"] == "Forecast" and by["PO-10001|balance"]["amount"] == 71.5,
          "future balance is Forecast and keeps the amount override")
    check(not by["PO-10001|balance"]["in_window"], "20 Nov is outside the 4-week window")
    check(by["PO-2|full"]["status"] == "Date TBC", "no ETA -> Date TBC")
    check(by["PO-3|deposit"]["status"] == "Paid" and by["PO-3|deposit"]["status_assumed"],
          "assumed-paid deposit")
    check(by["M1"]["status"] == "For Payment" and by["M1"]["state"] == "WA",
          "manual line with no due date goes in this request and inherits the PO state")
    check(by["M1"]["this_request"], "manual line is in this request batch")


def test_ledger():
    with tempfile.TemporaryDirectory() as tmp:
        L.DB = Path(tmp) / "t.db"
        L.set_line_state("PO-1|full", {"status": "Paid", "paid_date": "2026-09-25"}, "t")
        check(L.all_line_states()["PO-1|full"]["status"] == "Paid", "line state stored")
        L.set_line_state("PO-1|full", {"status": ""}, "t")
        check(L.all_line_states()["PO-1|full"]["status"] is None, "blank clears an override")
        for bad in ({"paid_date": "31/12/2026"}, {"status": "Maybe"}, {"amount": "abc"}):
            try:
                L.set_line_state("PO-1|full", bad, "t")
                check(False, f"invalid value accepted: {bad}")
            except ValueError:
                pass
        mid = L.add_manual({"pay_type": "Freight Payment", "amount": "1,200"}, "t")
        m = L.all_manual()[0]
        check(m["currency"] == "USD" and m["amount"] == 1200.0, "freight defaults to USD, amount parsed")
        check(L.add_manual({"pay_type": "Cartage", "amount": 1}, "t", seed_key="k") is not None and
              L.add_manual({"pay_type": "Cartage", "amount": 1}, "t", seed_key="k") is None,
              "seed key imports a line once")
        L.delete_manual(mid, "t")
        check(all(x["id"] != mid for x in L.all_manual()), "deleted line hidden")
        check(len(L.recent_audit()) >= 5, "every change is audited")
    L.DB = HERE / "ledger.db"


def test_data():
    cache_p, payload_p = HERE / "payment_cache.json", HERE / "report_payload.json"
    if not cache_p.exists() or not payload_p.exists():
        print("  (no cache/payload yet - data checks skipped)")
        return
    cache = json.loads(cache_p.read_text(encoding="utf-8"))
    payload = json.loads(payload_p.read_text(encoding="utf-8"))
    auto = payload["auto_lines"]
    bulk_listed = {ln["po"] for ln in auto if ln.get("bulk")}
    open_pos = {r for r, p in cache["pos"].items()
                if p["is_open"] and p["stage_raw"] not in P.BULK_STAGES} | bulk_listed
    listed = {ln["po"] for ln in auto}
    check(open_pos == listed, f"every open PO has payment lines ({len(open_pos ^ listed)} differ)")
    by_po: dict[str, float] = {}
    for ln in auto:
        by_po[ln["po"]] = by_po.get(ln["po"], 0) + ln["pct"]
        check(ln["amount_auto"] >= 0, f"{ln['line_id']} negative amount")
        check(ln["state"] in ("NSW", "VIC", "QLD", "WA", "SA", ""), f"{ln['line_id']} bad state {ln['state']}")
        check(ln["segment"] in ("Retail", "Commercial"), f"{ln['line_id']} bad segment")
    check(all(abs(v - 100) < 1e-6 for po, v in by_po.items() if po not in bulk_listed),
          "each PO's milestones add to 100% (a bulk order carries only its deposit)")
    check(len({ln["line_id"] for ln in auto}) == len(auto), "line ids are unique")
    no_state = [ln["po"] for ln in auto if not ln["state"] and not ln.get("bulk")]
    check(not no_state, f"POs with no state: {no_state[:5]}")
    wm = cache["meta"].get("po_watermark")
    check(bool(wm), "incremental watermark is set")
    rows, _ = P.merge(auto, {}, [], dt.date.today(), 4, payload.get("po_index"), payload.get("branch_state"))
    check(all(r["request_date"] is None or dt.date.fromisoformat(r["request_date"]).weekday() in (1, 3)
              for r in rows), "every pay-run date is a Tuesday or Thursday")


def test_uat_rules():
    """Rules added after the 6 Oct 2026 UAT."""
    import update_cache as U
    check(U._iso("2026-10-24T17:30:00Z") == "2026-10-25", "Cin7 UTC timestamp -> Perth date")
    check(U._iso("2026-10-24T03:00:00Z") == "2026-10-24", "morning UTC stays the same Perth day")
    rule = T.parse_terms({"terms": "20% Deposit, 80% against copy of shipping docs", "deposit_pct": "20"})
    cal = {"deposit_rate": 0.33, "deposit_pos": 36, "anchor": "eta", "offset_days": -9, "n": 51}
    r2 = P._calibrated(rule, cal)
    check(r2["deposit_pct"] == 0 and r2["balance_anchor"] == "eta" and r2["balance_offset"] == -9,
          "rarely-used deposit dropped, timing from history")
    check(P._calibrated(rule, {"deposit_rate": 0.9, "deposit_pos": 10})["deposit_pct"] == 20,
          "a deposit the team still uses is kept")
    po = {"is_open": True, "stage_raw": "New", "is_approved": False, "supplier": "X", "po_date": "2026-10-01",
          "eta_cin7": "2026-12-01", "value_usd": 100.0, "usd_exact": True, "ship": {}, "state": "WA"}
    lines, st = P.auto_lines({"pos": {"PO-50001": po}}, [], {}, {}, {})
    rows, _ = P.merge(lines, {}, [], D("2026-10-05"), 4)
    check(all(r["status"] == "Awaiting Approval" and not r["in_window"] for r in rows),
          "unapproved New PO is listed but never requested")
    po2 = dict(po, stage_raw="On Its Way", is_approved=True)
    lines, _ = P.auto_lines({"pos": {"PO-50002": po2}}, [], {}, {}, {})
    rows, _ = P.merge(lines, {"PO-50002|balance": {"status": "Requested", "request_date": "2026-10-06",
                                                    "amount": 99.0}}, [], D("2026-10-05"), 4)
    full = [r for r in rows if r["kind"] == "full"][0]
    check(full["status"] == "Requested" and full["amount"] == 99.0,
          "a record on the old balance line carries over to the single full payment")


def test_uat_13oct():
    """Rules added after the 13 Oct 2026 UAT (Idelyn)."""
    base = {"source": "auto", "po": "PO-60001", "po_list": ["PO-60001"], "vendor": "X", "state": "WA",
            "segment": "Retail", "currency": "USD", "flags": [], "assumed_paid": None, "kind": "full",
            "line_id": "PO-60001|full", "amount_auto": 10.0, "due_auto": "2026-10-08"}
    rows, when = P.merge([base], {"PO-60001|full": {"request_date": "2026-10-06", "amount": 9.0}},
                         [], D("2026-10-14"), 4)
    r = rows[0]
    check(r["request_date"] == "2026-10-20" and r["carried_from"] == "2026-10-06" and r["status"] == "For Payment",
          "unpaid line planned for a past pay run moves to the next one")
    check(when["next_request"] == "2026-10-20" and when["next_request_aud"] == "2026-10-15",
          "next runs: Tuesday USD, Thursday AUD")
    rows, _ = P.merge([base], {"PO-60001|full": {"status": "Requested", "request_date": "2026-10-06"}},
                      [], D("2026-10-14"), 4)
    check(rows[0]["request_date"] == "2026-10-06", "a line already sent to Accounts keeps its date")
    man = [{"id": 7, "po_refs": "", "pay_type": "GST Payment", "currency": "AUD", "amount": 5.0,
            "due_date": None, "status": None}]
    rows, _ = P.merge([], {}, man, D("2026-10-14"), 4)
    check(rows[0]["request_date"] == "2026-10-15", "AUD charge with no due date -> Thursday run")
    check(not P._evidence({"finalised": 1}) and P._evidence({"status": "Paid"}),
          "a ticked Final box alone is not evidence of a deposit")
    rule = T.parse_terms(None)
    r2 = P._calibrated(rule, {"deposit_rate": 1.0, "deposit_pos": 5, "deposit_pct": 30})
    check(r2["deposit_pct"] == 30, "a deposit the supplier always takes is added from history")
    po = {"is_open": True, "stage_raw": "On Its Way", "is_approved": True, "supplier": "X",
          "po_date": "2026-08-01", "eta_cin7": "2026-10-31", "value_usd": 100.0, "usd_exact": True,
          "state": "WA", "ship": {"eta": "2026-10-24", "eta_shipper": "2026-10-17", "etd": "2026-09-10"}}
    lines, _ = P.auto_lines({"pos": {"PO-60002": po}}, [], {}, {}, {})
    check(lines[0]["eta_port"] == "2026-10-24",
          "forwarder ETA follows a Cin7 ETA change made before the Shipment Schedule re-reads it")
    po["eta_cin7"] = "2026-10-25"
    lines, _ = P.auto_lines({"pos": {"PO-60002": po}}, [], {}, {}, {})
    check(lines[0]["eta_port"] == "2026-10-17", "a one-day rounding difference is ignored")


def test_oct15():
    """ACC remarks, payment archive, OneDrive backup layout (15 Oct 2026)."""
    import build_report as BR
    with tempfile.TemporaryDirectory() as tmp:
        L.DB = Path(tmp) / "t.db"
        rm = L.get_remarks()
        check("CSP - 14 days - prior to arrival" in rm["acc"]["text"] and "FDF - 45 days" in rm["acc"]["text"],
              "ACC block starts with the two points from the workbook")
        check(rm["sc"]["text"].startswith("K-Life"), "SC block starts with the known SC arrangements")
        saved = L.set_remark("acc", "Line one\r\nLine two", "acc@test")
        check(saved["text"] == "Line one\nLine two" and saved["updated_by"] == "acc@test", "remarks save with author")
        for bad in (("xx", "t"), ("sc", "x" * (L.REMARK_MAX + 1))):
            try:
                L.set_remark(*bad, "t")
                check(False, f"bad remark accepted: {bad[0]}")
            except ValueError:
                pass
        ln = {"line_id": "PO-70001|full", "po": "PO-70001", "kind": "full", "supplier": "X", "vendor": "X",
              "currency": "USD", "amount_auto": 10.0, "due_auto": "2026-10-01", "flags": [], "stage": "On Its Way"}
        check(BR.is_payment_record({"status": "Paid"}) and not BR.is_payment_record({"finalised": 1})
              and not BR.is_payment_record({"amount": 5.0}), "only paid / requested / written-off lines are archived")
        check(BR.effective_state(ln, {"PO-70001|balance": {"status": "Paid"}}) == {"status": "Paid"},
              "a full line reads its old balance record")
        L.archive_lines([ln])
        snaps = L.archived_snapshots()
        closed = BR.closed_archive(snaps, [], {"pos": {"PO-70001": {"stage": "Received"}}})
        check(len(closed) == 1 and closed[0]["archived"] and closed[0]["stage"] == "Received",
              "a recorded line outlives its PO once Cin7 closes it")
        check(BR.closed_archive(snaps, [dict(ln)], {}) == [], "while the PO is open the live line is used")
        rows, _ = P.merge(closed, {"PO-70001|full": {"status": "Paid", "paid_date": "2026-10-13"}}, [],
                          D("2026-10-15"), 4)
        check(rows[0]["status"] == "Paid" and rows[0]["paid_date"] == "2026-10-13", "archived line keeps its record")
        rows[0]["not_final"] = False
        path = BR.write_register_xlsx(rows, {"_remarks": L.get_remarks(), "_terms_table": []},
                                      Path(tmp) / "reg.xlsx")
        import openpyxl
        wb = openpyxl.load_workbook(path)
        hdr = [c.value for c in wb["updated"][1]]
        check(wb.sheetnames[:2] == ["updated", "App register"] and hdr[11] == "Payment Status/Type"
              and hdr[8] == "ETA\n(Forwarder)" and hdr[23] == "%" and len(hdr) == 26
              and hdr[24:] == ["Paid Currency", "Paid Amount (AUD)"],
              "backup uses the team's register layout, Accounts' two columns after it")
        check(wb["updated"]["L2"].value == "Paid" and "Special Arrangements" in wb.sheetnames,
              "status wording and remarks sheet in the backup")
        wb.close()
    L.DB = HERE / "ledger.db"
    check(BR.status_type({"source": "auto", "kind": "deposit", "pct": 30, "status": "Paid"}) == "30% deposit paid"
          and BR.status_type({"source": "auto", "kind": "balance", "status": "For Payment"}) == "Balance for payment"
          and BR.status_type({"source": "manual", "pay_type": "GST Payment", "status": "Requested"}) == "GST Payment",
          "Payment Status/Type follows the team's wording")


def test_uat_16oct():
    """Rules added after the 16 Oct 2026 UAT (Idelyn): slipped plans, draft POs paid
    in advance, a deposit typed on a full payment."""
    import seed_from_reference as SR
    base = {"source": "auto", "po": "PO-80001", "po_list": ["PO-80001"], "vendor": "X", "state": "VIC",
            "segment": "Retail", "currency": "USD", "flags": [], "assumed_paid": None, "kind": "full",
            "line_id": "PO-80001|full", "amount_auto": 10.0, "due_auto": "2026-11-17"}
    plan = {"request_date": "2026-10-20", "plan_due": "2026-10-20"}
    rows, _ = P.merge([base], {"PO-80001|full": plan}, [], D("2026-10-17"), 4)
    r = rows[0]
    check(r["request_date"] == "2026-11-17" and r["replanned_from"] == "2026-10-20" and r["status"] == "Forecast",
          "a planned date moves with the PO when its due date slips four weeks")
    check(any("moved +28 days" in f for f in r["flags"]) and base["flags"] == [],
          "the move is explained on the line (and the payload line is not touched)")
    rows, _ = P.merge([dict(base, due_auto="2026-10-25")], {"PO-80001|full": plan}, [], D("2026-10-17"), 4)
    check(rows[0]["request_date"] == "2026-10-20" and not rows[0].get("replanned_from"),
          "a move of a week or less keeps the plan")
    rows, _ = P.merge([base], {"PO-80001|full": dict(plan, plan_due=None)}, [], D("2026-10-17"), 4)
    check(rows[0]["request_date"] == "2026-10-20", "a plan earlier than the due date on purpose stays")
    rows, _ = P.merge([base], {"PO-80001|full": dict(plan, status="Requested")}, [], D("2026-10-17"), 4)
    check(rows[0]["request_date"] == "2026-10-20", "a line already sent to Accounts keeps its date")
    early = {"request_date": "2026-11-03", "plan_due": "2026-11-10"}
    rows, _ = P.merge([dict(base, due_auto="2026-11-01")], {"PO-80001|full": early}, [], D("2026-10-26"), 4)
    check(rows[0]["request_date"] == "2026-11-03" and any("earlier" in f for f in rows[0]["flags"]),
          "a PO that comes forward keeps its planned date and is flagged (Accounts may have scheduled it)")
    dep = dict(base, line_id="PO-80001|deposit", kind="deposit", due_anchor="order", due_auto="2026-10-23")
    rows, _ = P.merge([dep], {"PO-80001|deposit": {"request_date": "2026-10-20", "plan_due": "2026-10-06"}},
                      [], D("2026-10-17"), 4)
    check(rows[0]["request_date"] == "2026-10-20" and not rows[0].get("replanned_from"),
          "a deposit keeps its plan when approval moves the PO date (VOG PO-140079)")

    gc = {"Qingdao Great Century Plastic Co.,Ltd": {"deposit_pct": 0, "anchor": "order", "offset_days": 0,
                                                    "rule": "100% T/T in advance"}}
    terms = [{"supplier_name": "Qingdao Great Century Plastic Co.,Ltd", "terms": "50% Deposit, 50% balance "
              "based on copy of B/L", "deposit_pct": "50", "basis": "B/L copy"}]
    po = {"is_open": True, "stage_raw": "NEW", "is_approved": False, "po_date": "2026-10-10",
          "supplier": "Qingdao Great Century Plastic Co.,Ltd ", "value_usd": 804.0, "usd_exact": True,
          "ship": {}, "state": "NSW"}
    lines, _ = P.auto_lines({"pos": {"PO-141290": po}}, terms, {"special_arrangements": gc}, {}, {})
    check(len(lines) == 1 and lines[0]["kind"] == "full" and lines[0]["due_auto"] == "2026-10-10",
          "Great Century: one payment, due on the PO date (paid in advance)")
    rows, _ = P.merge(lines, {}, [], D("2026-10-17"), 4)
    check(rows[0]["status"] == "Awaiting Approval", "a draft PO with nothing scheduled stays Awaiting Approval")
    rows, _ = P.merge(lines, {"PO-141290|full": {"request_date": "2026-10-20", "amount": 804.0}}, [],
                      D("2026-10-17"), 4)
    check(rows[0]["status"] == "For Payment" and rows[0]["this_request"]
          and any("draft PO" in f for f in rows[0]["flags"]),
          "a request date set on a draft PO puts it in the pay run, flagged")
    rows, _ = P.merge(lines, {"PO-141290|full": {"amount": 3102.4}}, [], D("2026-10-17"), 4)
    check(rows[0]["status"] == "Awaiting Approval", "an amount alone does not schedule a draft PO")
    rule = T.parse_terms({"terms": "100% T/T in advance", "basis": "Prepaid"})
    check(P._calibrated(rule, {"anchor": "etd", "offset_days": -28, "n": 12})["balance_anchor"] == "order",
          "payment history does not move an in-advance term")

    check(P.deposit_pct_of("Deposit 30%") == 30 and P.deposit_pct_of("Deposit") == 0
          and P.deposit_pct_of("Full payment 100%") is None, "deposit percentage read from a payment type")
    hl = {"is_open": True, "stage_raw": "Production-Allocated", "is_approved": True, "po_date": "2026-10-15",
          "supplier": "HEALTHLAND SAUNA EQUIPMENT CO.,LTD", "value_usd": 13312.0, "usd_exact": True,
          "eta_cin7": "2027-01-20", "ship": {}, "state": "VIC"}
    st = {"PO-142054|deposit": {"pay_type": "Deposit 30%", "amount": 4243.0, "request_date": "2026-10-20",
                                "invoice_no": "HL261015-L-2"}}
    lines, _ = P.auto_lines({"pos": {"PO-142054": hl}}, [], {}, {}, st)
    by = {ln["kind"]: ln for ln in lines}
    check(set(by) == {"deposit", "balance"} and by["deposit"]["pct"] == 30 and by["deposit"]["due_auto"] == "2026-10-21",
          "a deposit typed on a full payment splits the PO; the deposit falls due after approval")
    rows, _ = P.merge(lines, st, [], D("2026-10-17"), 4)
    rb = {r["kind"]: r for r in rows}
    check(rb["deposit"]["amount"] == 4243.0 and rb["deposit"]["this_request"]
          and rb["balance"]["amount"] == 9069.0 and rb["balance"]["status"] == "Forecast",
          "the balance is the PO value less the deposit, and is not lost")
    lines, _ = P.auto_lines({"pos": {"PO-142054": hl}}, [], {}, {}, {"PO-142054|deposit": {"pay_type": "Deposit 30%"}})
    check({ln["kind"] for ln in lines} == {"deposit", "balance"}, "the deposit type alone is enough to split")

    with tempfile.TemporaryDirectory() as tmp:
        L.DB = Path(tmp) / "t.db"
        L.set_line_state("PO-142054|full", {"pay_type": "Deposit 30%", "amount": 4243}, "rizqi@test")
        check(L.split_full_line("PO-142054", "rizqi@test") and "PO-142054|full" not in L.all_line_states()
              and L.all_line_states()["PO-142054|deposit"]["amount"] == 4243.0,
              "the full line's record moves to the deposit line")
        check(not L.split_full_line("PO-142054", "t"), "a PO split once is not split again")
        with L.connect() as con:
            human = SR._human_fields(con, "PO-142054|deposit")
        check({"amount", "pay_type"} <= human, "the moved fields stay the person's (a sync keeps them)")
        L.set_line_state("PO-9|balance", {"request_date": "2026-10-20"}, "seed: payment request workbook")
        L.split_full_line("PO-9", "t")
        check("PO-9|deposit" in L.all_line_states(), "the old balance record a full line reads moves too")
        L.set_line_state("PO-8|full", {"request_date": "2026-10-20", "plan_due": "2026-10-21"}, "t")
        L.set_line_state("PO-7|full", {"request_date": "2026-10-20"}, "t")
        check(L.backfill_plan_due({"PO-8|full": "2026-12-01", "PO-7|full": "2026-10-22"}) == 1
              and L.all_line_states()["PO-8|full"]["plan_due"] == "2026-10-21"
              and L.all_line_states()["PO-7|full"]["plan_due"] == "2026-10-22",
              "backfill gives a due date only to plans that have none")
        L.set_line_state("PO-6|full", {"request_date": "2026-10-20", "plan_due": "2026-10-20"},
                         "seed: payment request workbook")
        SR.sync({"PO-6|full": {"request_date": "2026-10-20", "plan_due": "2026-11-17", "_row": 1}}, [], "t")
        check(L.all_line_states()["PO-6|full"]["plan_due"] == "2026-10-20",
              "the same date in a newer workbook does not re-anchor a plan")
        SR.sync({"PO-6|full": {"request_date": "2026-11-17", "plan_due": "2026-11-17", "_row": 1}}, [], "t")
        check(L.all_line_states()["PO-6|full"]["plan_due"] == "2026-11-17",
              "a new date in the workbook is a new plan")
    L.DB = HERE / "ledger.db"


def test_oct26():
    """Accounts' paid currency and paid amount in AUD (26 Oct 2026)."""
    import seed_from_reference as SR
    check(L.clean("paid_currency", "cny") == "CNY" and L.clean("paid_currency", "") is None,
          "paid currency accepts CNY, blank = automatic")
    try:
        L.clean("paid_currency", "XYZ")
        check(False, "an unknown paid currency is accepted")
    except ValueError:
        pass
    check(L.clean("paid_amount_aud", "3,250.5") == 3250.5, "paid amount (AUD) parsed")
    check(SR._paid_currency("RMB") == "CNY" and SR._paid_currency("Euro") == "EUR"
          and SR._paid_currency("bitcoin") is None, "workbook paid currencies are tidied, unknown ones ignored")
    base = {"source": "auto", "po": "PO-90001", "po_list": ["PO-90001"], "vendor": "X", "state": "WA",
            "segment": "Retail", "currency": "USD", "flags": [], "assumed_paid": None, "kind": "full",
            "line_id": "PO-90001|full", "amount_auto": 10.0, "due_auto": "2026-10-20"}
    rows, _ = P.merge([base], {"PO-90001|full": {"status": "Paid", "paid_amount": 10.0}}, [], D("2026-10-26"), 4)
    r = rows[0]
    check(r["paid_currency"] == "USD" and r["paid_currency_default"] and r["paid_amount_aud"] is None,
          "paid in the line's currency by default; the AUD cost is left to Accounts")
    rows, _ = P.merge([base], {"PO-90001|full": {"status": "Paid", "paid_amount": 10.0, "paid_currency": "CNY",
                                                 "paid_amount_aud": 2.1}}, [], D("2026-10-26"), 4)
    check((rows[0]["paid_currency"], rows[0]["paid_amount_aud"], rows[0]["paid_aud_default"]) == ("CNY", 2.1, False),
          "what Accounts types wins")
    man = [{"id": 9, "po_refs": "PO-90001", "pay_type": "Customs Brokerage", "currency": "AUD", "amount": 800.0,
            "due_date": None, "status": "Paid", "paid_amount": 800.0}]
    rows, _ = P.merge([], {}, man, D("2026-10-26"), 4)
    check(rows[0]["paid_currency"] == "AUD" and rows[0]["paid_amount_aud"] == 800.0 and rows[0]["paid_aud_default"],
          "an AUD charge paid in AUD fills the AUD column by itself")
    rows, _ = P.merge([base], {}, [], D("2026-10-26"), 4)
    check(rows[0]["paid_currency"] is None and rows[0]["paid_amount_aud"] is None, "nothing paid, nothing shown")


def test_oct27():
    """ACC check of the 27 Oct pay run (Arlene): K-Life deposits covered by the bulk
    order, and timing from scattered history (Rizhao Pioneer)."""
    import build_report as BR
    kl = "Xiamen K-Life Sport Co., Ltd"
    ov = {"bulk_deposit_suppliers": [kl], "term_aliases": {kl: "KLife"}}
    terms = [{"supplier_name": "KLife", "terms": "30% deposit, 70% balance before shipment", "deposit_pct": "30"}]
    ship = {"is_open": True, "stage_raw": "Production-Allocated", "is_approved": True, "supplier": kl,
            "po_date": "2026-10-20", "usd_exact": True, "eta_cin7": "2027-01-09", "ship": {}, "state": "WA"}
    cache = {"pos": {
        "PO-135014": dict(ship, stage_raw="Production-China Stock", po_date="2026-08-21", value_usd=57090.0),
        "PO-142640": dict(ship, value_usd=15858.0), "PO-142641": dict(ship, value_usd=22770.0),
        "PO-142642": dict(ship, value_usd=10470.0),
        "PO-142999": dict(ship, po_date="2026-10-25", value_usd=10000.0)}}
    paid = {"PO-135014|deposit": {"status": "Paid", "paid_date": "2026-10-13", "paid_amount": 17127.0}}
    lines, stats = P.auto_lines(cache, terms, ov, {}, paid)
    dep = {ln["po"]: ln for ln in lines if ln["kind"] == "deposit"}
    check(all(dep[p].get("allocated_from") == "PO-135014" and not dep[p]["review"]
              for p in ("PO-142640", "PO-142641", "PO-142642")) and stats["bulk_allocated"] == 3,
          "K-Life deposits covered by the paid bulk deposit are allocated (14,729.40 of 17,127)")
    check(not dep["PO-142999"].get("allocated_from") and dep["PO-142999"]["review"]
          and any("not enough" in f for f in dep["PO-142999"]["flags"]),
          "a deposit the bulk deposit no longer covers stays CHECK (USD 2,397.60 left < 3,000)")
    rows, _ = P.merge(lines, paid, [], D("2026-10-27"), 4)
    by = {r["line_id"]: r for r in rows}
    r = by["PO-142641|deposit"]
    check(r["status"] == "Not Payable" and r["status_assumed"] and not r["this_request"]
          and BR.status_type(r) == "Allocated from PO-135014",
          "an allocated deposit is not in the pay run and reads 'Allocated from PO-135014'")
    check(by["PO-135014|deposit"]["status"] == "Paid" and by["PO-142641|balance"]["status"] in ("Forecast", "For Payment"),
          "the bulk deposit and the shipment POs' balances are untouched")
    lines, _ = P.auto_lines(cache, terms, ov, {}, {**paid, "PO-142640|deposit": {"amount": 4757.4}})
    check(not [ln for ln in lines if ln["line_id"] == "PO-142640|deposit"][0].get("allocated_from"),
          "a deposit someone recorded an amount on is left to them")
    lines, _ = P.auto_lines(cache, terms, ov, {}, {})
    check(not any(ln.get("allocated_from") for ln in lines), "nothing is allocated before the bulk deposit is paid")

    cal = {"deposit_rate": 0.5, "deposit_pos": 4, "deposit_pct": 30.0, "anchor": "etd", "offset_days": -11,
           "n": 12, "mad": 22.5}
    r2 = P._calibrated(T.parse_terms(None), cal)
    check((r2["deposit_pct"], r2["balance_anchor"], r2["balance_offset"]) == (30.0, "eta", -7)
          and "scatter" in r2["rule"], "scattered history (+-22.5 days) does not set the timing; the default does")
    r3 = P._calibrated(T.parse_terms(None), dict(cal, mad=3))
    check((r3["balance_anchor"], r3["balance_offset"]) == ("etd", -11), "consistent history still sets the timing")


def main() -> int:
    for fn in (test_request_days, test_po_list, test_terms, test_merge, test_ledger, test_uat_rules, test_uat_13oct,
               test_oct15, test_uat_16oct, test_oct26, test_oct27, test_data):
        print(fn.__name__)
        fn()
    print(f"\n{'FAILED' if FAILS else 'OK'}: {len(FAILS)} failure(s)")
    return 1 if FAILS else 0


if __name__ == "__main__":
    raise SystemExit(main())
